#!/usr/bin/env python3
"""Run a deterministic positive control through the official Toolathlon runtime."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from research_ops.openreward import (
    RemoteORSBackend,
    RemoteORSConfig,
    build_default_environments,
    create_openreward_app,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PINNED_REVISION = "9c82859d05bc1f0376150d749cd2e3d43471d583"
TASK_NAME = "canvas-assignment-stats"
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "data"
    / "openreward_tooluse_canary_20260720"
    / "canvas-assignment-stats_positive_evidence.json"
)
_INFRASTRUCTURE_FAILURES = (
    "Traceback (most recent call last)",
    "Evaluation error:",
    "connection to server",
    "could not connect to server",
    "FATAL:",
    "is not running",
)


def _sse_output(response: object) -> dict[str, Any]:
    text = str(getattr(response, "text", ""))
    for line in text.splitlines():
        if not line.startswith("data: "):
            continue
        try:
            payload = json.loads(line[6:])
        except json.JSONDecodeError:
            continue
        if payload.get("ok") is True:
            return dict(payload["output"])
        if payload.get("ok") is False:
            raise RuntimeError(f"ToolUse ORS action failed: {payload.get('error')}")
    raise RuntimeError("ToolUse ORS response omitted a terminal SSE payload")


def _output_text(output: dict[str, Any]) -> str:
    blocks = output.get("blocks") or []
    if not isinstance(blocks, list):
        return ""
    return "\n".join(
        str(block.get("text") or "")
        for block in blocks
        if isinstance(block, dict) and block.get("type") == "text"
    )


def _call(
    client: TestClient,
    headers: dict[str, str],
    name: str,
    inputs: dict[str, Any],
) -> dict[str, Any]:
    return _sse_output(
        client.post(
            "/tooluse/call",
            headers=headers,
            json={"name": name, "input": inputs},
        )
    )


def _mcp_json(
    client: TestClient,
    headers: dict[str, str],
    name: str,
    arguments: dict[str, Any],
) -> Any:
    output = _call(
        client,
        headers,
        "call_tool",
        {"name": name, "arguments": arguments},
    )
    text = _output_text(output).strip()
    if text.casefold().startswith("error:"):
        raise RuntimeError(f"{name} failed: {text}")
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"{name} returned non-JSON output") from exc


def _assignment_rows(
    client: TestClient,
    headers: dict[str, str],
) -> list[dict[str, Any]]:
    courses = _mcp_json(client, headers, "canvas.canvas_list_courses", {})
    if not isinstance(courses, list) or not courses:
        raise RuntimeError("Canvas returned no courses")
    rows: list[dict[str, Any]] = []
    for course in courses:
        if not isinstance(course, dict):
            raise RuntimeError("Canvas returned an invalid course record")
        course_id = course.get("id")
        course_code = str(course.get("course_code") or "").strip()
        if course_id is None or not course_code:
            raise RuntimeError("Canvas course omitted its id or course code")
        assignments = _mcp_json(
            client,
            headers,
            "canvas.canvas_list_assignments",
            {"course_id": course_id},
        )
        if not isinstance(assignments, list):
            raise RuntimeError("Canvas returned invalid assignment data")
        names = [
            str(item.get("name") or "").strip()
            for item in assignments
            if isinstance(item, dict)
        ]
        points = [
            float(item["points_possible"])
            for item in assignments
            if isinstance(item, dict) and item.get("points_possible") is not None
        ]
        rows.append(
            {
                "course_code": course_code,
                "total_assignments": len(assignments),
                "avg_points": round(sum(points) / len(points), 1) if points else 0.0,
                "tma_count": sum(name.casefold().startswith("tma") for name in names),
                "cma_count": sum(name.casefold().startswith("cma") for name in names),
            }
        )
    return sorted(rows, key=lambda item: item["course_code"])


def _write_artifacts(
    client: TestClient,
    headers: dict[str, str],
    rows: list[dict[str, Any]],
) -> tuple[str, str]:
    total = sum(int(item["total_assignments"]) for item in rows)
    total_tmas = sum(int(item["tma_count"]) for item in rows)
    total_cmas = sum(int(item["cma_count"]) for item in rows)
    most_assignments = max(rows, key=lambda item: item["total_assignments"])
    serialized = json.dumps(rows, separators=(",", ":"), sort_keys=True)
    code = f'''import hashlib
import json
from openpyxl import Workbook
from docx import Document

rows = json.loads({serialized!r})
workbook = Workbook()
stats = workbook.active
stats.title = "Assignment Stats"
stats.append(["Course_Code", "Total_Assignments", "Avg_Points", "TMA_Count", "CMA_Count"])
for row in rows:
    stats.append([
        row["course_code"],
        row["total_assignments"],
        row["avg_points"],
        row["tma_count"],
        row["cma_count"],
    ])
summary = workbook.create_sheet("Summary")
summary.append(["Metric", "Value"])
summary.append(["Total_Assignments", {total}])
summary.append(["Total_TMAs", {total_tmas}])
summary.append(["Total_CMAs", {total_cmas}])
summary.append(["Course_Most_Assignments", {most_assignments["course_code"]!r}])
xlsx_path = "Canvas_Assignment_Stats.xlsx"
workbook.save(xlsx_path)

document = Document()
document.add_heading("Assignment Overview", level=1)
document.add_paragraph(
    "This assignment overview covers " + str(len(rows)) + " Canvas courses and "
    + str({total}) + " assignments. It includes " + str({total_tmas})
    + " Tutor Marked Assessments and " + str({total_cmas})
    + " Computer Marked Assessments."
)
docx_path = "Assignment_Overview.docx"
document.save(docx_path)
for label, path in (("xlsx", xlsx_path), ("docx", docx_path)):
    with open(path, "rb") as handle:
        print(label + "_sha256=" + hashlib.sha256(handle.read()).hexdigest())
'''
    output = _call(client, headers, "python_execute", {"code": code})
    text = _output_text(output)
    hashes = dict(line.split("=", 1) for line in text.splitlines() if "_sha256=" in line)
    if len(hashes.get("xlsx_sha256", "")) != 64 or len(hashes.get("docx_sha256", "")) != 64:
        raise RuntimeError("workspace did not produce the expected artifacts")
    return hashes["xlsx_sha256"], hashes["docx_sha256"]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--service-url-env", default="ORS_TOOLUSE_SERVICE_URL")
    parser.add_argument("--service-token-env")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    service_url = str(os.environ.get(args.service_url_env, "")).strip()
    if not service_url:
        raise SystemExit(f"missing Toolathlon service reference: {args.service_url_env}")
    auth = (
        {"Authorization": args.service_token_env}
        if args.service_token_env
        else {}
    )
    backend = RemoteORSBackend(
        RemoteORSConfig(
            environment="tooluse",
            remote_environment="toolathlongym",
            base_url_env=args.service_url_env,
            auth_header_envs=auth,
            timeout_seconds=300,
        )
    )
    task = {
        "task_id": TASK_NAME,
        "task_name": TASK_NAME,
        "query": "Use the prompt supplied by the official Toolathlon environment.",
        "source": {
            "repository": "EnvCommons/toolathlon-gym",
            "revision": PINNED_REVISION,
        },
    }
    app = create_openreward_app(
        build_default_environments(
            backends={"tooluse": backend},
            task_catalogs={"tooluse": {"train": [task]}},
        )
    )
    started = datetime.now(UTC)
    with TestClient(app) as client:
        sid = client.post("/create_session").json()["sid"]
        headers = {"X-Session-ID": sid}
        try:
            created = client.post(
                "/create",
                headers=headers,
                json={"env_name": "tooluse", "split": "train", "index": 0},
            )
            created.raise_for_status()
            rows = _assignment_rows(client, headers)
            xlsx_sha256, docx_sha256 = _write_artifacts(client, headers, rows)
            total = sum(int(item["total_assignments"]) for item in rows)
            total_tmas = sum(int(item["tma_count"]) for item in rows)
            total_cmas = sum(int(item["cma_count"]) for item in rows)
            terminal = _call(client, headers, "claim_done", {})
            evaluator_text = _output_text(terminal)
        finally:
            client.post("/delete", headers=headers)

    if terminal.get("reward") != 1.0 or terminal.get("finished") is not True:
        raise RuntimeError(f"official Toolathlon positive control failed: {evaluator_text}")
    if any(marker.casefold() in evaluator_text.casefold() for marker in _INFRASTRUCTURE_FAILURES):
        raise RuntimeError("Toolathlon evaluator reported an infrastructure failure")
    evidence = {
        "version": 1,
        "created_at": datetime.now(UTC).isoformat(),
        "task_id": TASK_NAME,
        "source_revision": PINNED_REVISION,
        "control": "deterministic_official_mcp_positive_control",
        "policy_driven": False,
        "reward": terminal.get("reward"),
        "finished": terminal.get("finished"),
        "course_count": len(rows),
        "assignment_count": total,
        "tma_count": total_tmas,
        "cma_count": total_cmas,
        "xlsx_sha256": xlsx_sha256,
        "docx_sha256": docx_sha256,
        "evaluator_output_sha256": hashlib.sha256(evaluator_text.encode()).hexdigest(),
        "duration_seconds": (datetime.now(UTC) - started).total_seconds(),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(evidence, indent=2, sort_keys=True) + "\n")
    print(json.dumps(evidence, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
